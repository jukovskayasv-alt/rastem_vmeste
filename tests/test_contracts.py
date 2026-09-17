import copy
from pathlib import Path
import tempfile
import unittest

from docx import Document

from skif_agents.contracts import create_contract


def deal():
    return {
        "demo": True,
        "cadastral": "DEMO-ONLY",
        "address": "Вымышленный район, демонстрационный участок",
        "area_m2": "600",
        "category": "Земли сельскохозяйственного назначения",
        "vri": "Ведение личного подсобного хозяйства на полевых участках",
        "land_regime": "field_lph",
        "construction_allowed": False,
        "seller_name": "Демонстрационный продавец",
        "buyer_name": "Демонстрационный покупатель",
        "price": "123456.78",
        "contract_date": "2026-09-17",
        "deadline": "2026-10-17",
        "transfer_days": 10,
        "payment_method": "Безналичный перевод по согласованным реквизитам",
        "payment_due_date": "2026-10-17",
        "installment": {"down_payment": "23456.78", "months": 3, "start_date": "2026-10-17"},
    }


def document_text(document):
    texts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        texts.extend(cell.text for row in table.rows for cell in row.cells)
    return "\n".join(texts)


class ContractTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.output = Path(self.temporary.name) / "draft.docx"

    def test_preliminary_is_unsigned_draft_and_uses_supplied_terms(self):
        result = create_contract("preliminary", deal(), self.output)
        self.assertEqual(result, self.output)
        self.assertTrue(result.is_file())
        document = Document(result)
        text = document_text(document)
        self.assertIn("ЧЕРНОВИК", text)
        self.assertIn("ДЕМОНСТРАЦИЯ", text)
        self.assertIn("DEMO-ONLY", text)
        self.assertIn("17.10.2026", text)
        self.assertIn("123 456,78", text)
        self.assertIn("строительство запрещено", text)
        self.assertIn("Демонстрационный продавец", text)
        self.assertIn("Основной договор", text)
        self.assertIn("не устанавливает обязанность оплаты", text)
        self.assertNotIn("электронной подписью подписан", text)
        self.assertTrue(any("ЧЕРНОВИК" in p.text for p in document.sections[0].header.paragraphs))

    def test_main_appendix_has_complete_reconciled_installment_schedule(self):
        create_contract("main", deal(), self.output)
        self.assertTrue(self.output.exists())
        document = Document(self.output)
        table = document.tables[-1]
        self.assertEqual(len(table.rows), 5)
        values = [[cell.text for cell in row.cells] for row in table.rows]
        self.assertEqual(values[1], ["0", "17.10.2026", "23 456,78", "100 000,00"])
        self.assertEqual(values[-1], ["3", "17.01.2027", "33 333,34", "0,00"])
        text = document_text(document)
        self.assertIn("10 календарных дней", text)
        self.assertIn("Безналичный перевод", text)
        self.assertIn("регистрации перехода права", text)
        self.assertIn("Залог", text)
        self.assertIn("согласие супруга", text)

    def test_main_without_installment_uses_explicit_full_payment_date(self):
        data = deal()
        del data["installment"]
        create_contract("main", data, self.output)
        self.assertTrue(self.output.exists())
        text = document_text(Document(self.output))
        self.assertIn("17.10.2026", text)
        self.assertNotIn("График платежей", text)

    def test_missing_essential_fields_are_not_invented(self):
        for field in ["cadastral", "address", "area_m2", "category", "vri", "seller_name", "buyer_name",
                      "price", "contract_date", "deadline", "land_regime", "construction_allowed"]:
            data = deal()
            del data[field]
            with self.subTest(field=field):
                with self.assertRaises(ValueError):
                    create_contract("preliminary", data, self.output)
                self.assertFalse(self.output.exists())

    def test_missing_main_specific_terms_are_rejected(self):
        for field in ["transfer_days", "payment_method"]:
            data = deal()
            del data[field]
            with self.subTest(field=field):
                with self.assertRaises(ValueError):
                    create_contract("main", data, self.output)
        data = deal()
        del data["installment"]
        del data["payment_due_date"]
        with self.assertRaises(ValueError):
            create_contract("main", data, self.output)

    def test_construction_prohibition_cannot_be_silently_changed(self):
        for changes in [{"land_regime": "izhs"}, {"construction_allowed": True}, {"construction_allowed": "false"}]:
            with self.subTest(changes=changes):
                with self.assertRaises(ValueError):
                    create_contract("main", deal() | changes, self.output)

    def test_real_mode_requires_cadastral_and_demo_cannot_use_real_number(self):
        for changes in [{"demo": False}, {"demo": "true"}, {"cadastral": "50:31:0012345:2371"},
                        {"demo": False, "cadastral": "2371"}]:
            with self.subTest(changes=changes):
                with self.assertRaises(ValueError):
                    create_contract("main", deal() | changes, self.output)
        real = deal() | {"demo": False, "cadastral": "50:31:0012345:100"}
        create_contract("main", real, self.output)
        self.assertTrue(self.output.exists())
        self.assertNotIn("ДЕМОНСТРАЦИЯ", document_text(Document(self.output)))

    def test_bad_numeric_and_date_fields_are_rejected(self):
        for changes in [{"area_m2": "0"}, {"area_m2": "NaN"}, {"area_m2": "600.001"},
                        {"price": "0"}, {"price": "123.456"}, {"contract_date": "2026-02-30"},
                        {"deadline": "2026-09-16"}, {"transfer_days": True}, {"transfer_days": 0},
                        {"seller_name": ""}, {"buyer_name": "A\x00B"}]:
            with self.subTest(changes=changes):
                kind = "preliminary" if "deadline" in changes else "main"
                with self.assertRaises(ValueError):
                    create_contract(kind, deal() | changes, self.output)

    def test_installment_cannot_supply_conflicting_price_or_precontract_payments(self):
        for changes in [{"price": "999"}, {"start_date": "2026-09-01"}, {"months": 0}, {"interest_rate": "3"}]:
            data = copy.deepcopy(deal())
            data["installment"].update(changes)
            with self.subTest(changes=changes):
                with self.assertRaises(ValueError):
                    create_contract("main", data, self.output)

    def test_existing_draft_is_not_overwritten(self):
        self.output.write_bytes(b"existing document")
        with self.assertRaises(FileExistsError):
            create_contract("main", deal(), self.output)
        self.assertEqual(self.output.read_bytes(), b"existing document")

    def test_unknown_kind_and_wrong_extension_are_rejected(self):
        with self.assertRaises(ValueError):
            create_contract("final", deal(), self.output)
        with self.assertRaises(ValueError):
            create_contract("main", deal(), self.output.with_suffix(".pdf"))


if __name__ == "__main__":
    unittest.main()

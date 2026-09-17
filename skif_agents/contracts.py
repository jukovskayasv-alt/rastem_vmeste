"""Create visibly unsigned draft agreements for field-LPH land only.

Validation checks data shape and internal consistency, never legal clearance.
The caller supplies every party and land fact. See docs/contracts.md.
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .finance import _iso_date, _money_kopecks, _rub, build_schedule


_CADASTRAL = re.compile(r"[0-9]{2}:[0-9]{2}:[0-9]{6,7}:[0-9]{1,10}\Z")
_REQUIRED_TEXT = ("cadastral", "address", "area_m2", "category", "vri", "seller_name",
                  "buyer_name", "price", "contract_date")


def _text(data: dict, field: str) -> str:
    value = data.get(field)
    if not isinstance(value, str) or not value.strip() or len(value) > 2000:
        raise ValueError(f"{field}: требуется непустая строка длиной до 2000 символов")
    if any(ord(character) < 32 or 0xD800 <= ord(character) <= 0xDFFF for character in value):
        raise ValueError(f"{field}: недопустимые управляющие символы")
    return value.strip()


def _validate(kind: str, data: dict) -> tuple[dict, dict | None]:
    if kind not in ("preliminary", "main"):
        raise ValueError("kind: допустимы preliminary и main")
    if not isinstance(data, dict):
        raise ValueError("Нужна карточка сделки в виде объекта")
    clean = {field: _text(data, field) for field in _REQUIRED_TEXT}
    demo = data.get("demo", False)
    if type(demo) is not bool:
        raise ValueError("demo: нужно логическое значение")
    clean["demo"] = demo
    if demo:
        if clean["cadastral"] != "DEMO-ONLY":
            raise ValueError("Для демонстрации cadastral должен быть DEMO-ONLY")
    elif not _CADASTRAL.fullmatch(clean["cadastral"]):
        raise ValueError("cadastral: нужен полный кадастровый номер, например AA:BB:CCCCCCC:D")
    if data.get("land_regime") != "field_lph" or data.get("construction_allowed") is not False:
        raise ValueError("Шаблон требует land_regime=field_lph и construction_allowed=false")
    area = _money_kopecks(clean["area_m2"], "area_m2")
    price = _money_kopecks(clean["price"], "price")
    if area <= 0 or price <= 0:
        raise ValueError("Площадь и цена должны быть положительными")
    clean["area_m2"], clean["price"] = _rub(area), _rub(price)
    contract_date = _iso_date(clean["contract_date"], "contract_date")
    if kind == "preliminary":
        clean["deadline"] = _text(data, "deadline")
        if _iso_date(clean["deadline"], "deadline") <= contract_date:
            raise ValueError("deadline: срок основного договора должен быть позже даты предварительного")
    else:
        transfer_days = data.get("transfer_days")
        if type(transfer_days) is not int or not 1 <= transfer_days <= 365:
            raise ValueError("transfer_days: нужно целое число от 1 до 365")
        clean["transfer_days"] = transfer_days
        clean["payment_method"] = _text(data, "payment_method")

    schedule = None
    if "installment" in data:
        installment = data["installment"]
        if not isinstance(installment, dict) or set(installment) != {"down_payment", "months", "start_date"}:
            raise ValueError("installment: нужны только down_payment, months и start_date; проценты не поддерживаются")
        schedule = build_schedule(clean["price"], installment["down_payment"],
                                  installment["months"], installment["start_date"])
        earliest = _iso_date(clean["deadline"], "deadline") if kind == "preliminary" else contract_date
        if _iso_date(schedule["start_date"], "start_date") < earliest:
            raise ValueError("Дата начала графика не может предшествовать дате основного договора или сроку его заключения")
    elif kind == "main":
        clean["payment_due_date"] = _text(data, "payment_due_date")
        if _iso_date(clean["payment_due_date"], "payment_due_date") < contract_date:
            raise ValueError("payment_due_date: дата оплаты не может быть раньше даты договора")
    return clean, schedule


def _display_date(value: str) -> str:
    year, month, day = value.split("-")
    return f"{day}.{month}.{year}"


def _display_money(value: str) -> str:
    whole, fraction = value.split(".")
    return f"{int(whole):,}".replace(",", " ") + "," + fraction


def _setup_document() -> Document:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin, section.bottom_margin = Inches(0.8), Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.8)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Header", "Footer"):
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(11)
        for border in style.element.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)
    normal = document.styles["Normal"].paragraph_format
    normal.space_after = Pt(7)
    normal.line_spacing = 1.12
    normal.keep_together = True
    title = document.styles["Title"]
    title.font.size, title.font.bold = Pt(18), True
    title.paragraph_format.space_after = Pt(11)
    title.paragraph_format.keep_with_next = True
    heading = document.styles["Heading 1"]
    heading.font.size, heading.font.bold = Pt(12), True
    heading.paragraph_format.space_before = Pt(12)
    heading.paragraph_format.space_after = Pt(6)
    heading.paragraph_format.keep_with_next = True
    section.header.paragraphs[0].text = "СКИФ  |  ЧЕРНОВИК  |  НЕ ДЛЯ ПОДПИСАНИЯ"
    section.header.paragraphs[0].runs[0].font.size = Pt(9)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("ЧЕРНОВИК  ·  ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    for run in footer.runs:
        run.font.size = Pt(9)
    document.core_properties.author = "СКИФ"
    document.core_properties.subject = "Неподписанный проект для юридической проверки"
    return document


def _schedule_table(document: Document, schedule: dict) -> None:
    table = document.add_table(rows=1, cols=4)
    table.autofit = False
    widths = [Inches(0.5), Inches(1.4), Inches(2.4), Inches(2.4)]
    for column, width in zip(table.columns, widths):
        column.width = width
    for cell, label in zip(table.rows[0].cells, ("№", "Дата", "Платеж ₽", "Остаток ₽")):
        cell.text = label
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for item in schedule["rows"]:
        cells = table.add_row().cells
        for cell, value in zip(cells, (str(item["number"]), _display_date(item["date"]),
                                      _display_money(item["amount"]), _display_money(item["balance"]))):
            cell.text = value
    for index, row in enumerate(table.rows):
        keep = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(keep)
        for position, cell in enumerate(row.cells):
            cell.width = widths[position]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            properties = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for side in ("top", "left", "bottom", "right"):
                border = OxmlElement(f"w:{side}")
                for name, value in {"val": "single", "sz": "6", "color": "D9D9D9"}.items():
                    border.set(qn(f"w:{name}"), value)
                borders.append(border)
            properties.append(borders)
            margins = OxmlElement("w:tcMar")
            for side in ("top", "left", "bottom", "right"):
                margin = OxmlElement(f"w:{side}")
                margin.set(qn("w:w"), "90")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            properties.append(margins)
            if index == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "E7EEF5")
                properties.append(shading)
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if position < 2 else WD_ALIGN_PARAGRAPH.RIGHT
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                    run.bold = index == 0


def create_contract(kind: str, data: dict, output: Path) -> Path:
    """Create a new DOCX draft, rejecting missing terms and existing files.

    `kind` is preliminary or main. This function neither signs a contract,
    verifies title, uploads personal data, nor registers a transfer of rights.
    """
    output = Path(output)
    if output.suffix.lower() != ".docx":
        raise ValueError("Файл договора должен иметь расширение .docx")
    clean, schedule = _validate(kind, data)
    if output.exists():
        raise FileExistsError(f"Черновик уже существует: {output}")
    document = _setup_document()
    title = ("Предварительный договор купли продажи земельного участка" if kind == "preliminary"
             else "Договор купли продажи земельного участка")
    document.add_paragraph(title, style="Title")
    notice = document.add_paragraph()
    notice.add_run("ЧЕРНОВИК ДЛЯ ЮРИДИЧЕСКОЙ ПРОВЕРКИ").bold = True
    document.add_paragraph("Проект условий для обсуждения продавцом, покупателем и юристом. "
                           "Документ не подписан и не подтверждает право собственности, оплату или регистрацию. "
                           "До подписания необходимо проверить сведения и дополнить реквизиты сторон.")
    if clean["demo"]:
        paragraph = document.add_paragraph()
        paragraph.add_run("ДЕМОНСТРАЦИЯ. ").bold = True
        paragraph.add_run("Данные сторон, участка и цены вымышлены. DEMO-ONLY не является кадастровым номером. "
                          "Этот пример не относится к реальному участку и не является предложением продажи.")
    document.add_paragraph(f"Дата проекта: {_display_date(clean['contract_date'])}.")
    document.add_paragraph(f"Продавец: {clean['seller_name']}. Покупатель: {clean['buyer_name']}.")

    document.add_heading("Предмет и характеристики участка", level=1)
    if kind == "preliminary":
        document.add_paragraph("Стороны обязуются заключить основной договор купли-продажи указанного ниже "
                               "земельного участка на согласованных в настоящем проекте условиях после их проверки.")
    else:
        document.add_paragraph("Продавец обязуется передать в собственность Покупателя указанный ниже земельный "
                               "участок, а Покупатель обязуется принять его и уплатить согласованную цену.")
    document.add_paragraph(f"Кадастровый номер: {clean['cadastral']}. Адрес или описание местоположения: {clean['address']}.")
    document.add_paragraph(f"Площадь: {_display_money(clean['area_m2'])} м². Категория земель: {clean['category']}. "
                           f"Вид разрешенного использования: {clean['vri']}.")
    document.add_paragraph("Основание права продавца, запись ЕГРН и дата актуальной выписки: "
                           "[ВНЕСТИ ПО ПОДТВЕРЖДАЮЩИМ ДОКУМЕНТАМ]. Наличие и содержание обременений "
                           "и прав третьих лиц: [ПРОВЕРИТЬ И ОПИСАТЬ].")

    document.add_heading("Цена и расчеты", level=1)
    document.add_paragraph(f"Цена всего участка составляет {_display_money(clean['price'])} руб. "
                           "Цена относится только к участку, описанному в этом проекте, и требует согласования сторонами.")
    if kind == "preliminary":
        document.add_paragraph(f"Основной договор стороны обязуются заключить не позднее {_display_date(clean['deadline'])}. "
                               "Способ обмена предложением о его заключении и адреса для уведомлений: [СОГЛАСОВАТЬ].")
        document.add_paragraph("Предварительный договор не устанавливает обязанность оплаты участка до заключения "
                               "основного договора. Аванс, задаток и обеспечительный платеж этим проектом не предусмотрены.")
        if schedule:
            document.add_paragraph("Приложение 1 содержит предлагаемый график расчетов по будущему основному договору. "
                                   "Его даты и способ расчетов необходимо согласовать при заключении основного договора; "
                                   "само приложение не является требованием об оплате.")
    else:
        document.add_paragraph(f"Способ расчетов: {clean['payment_method']}. Реквизиты получателя и документы, "
                               "подтверждающие исполнение платежа: [СОГЛАСОВАТЬ И ПРОВЕРИТЬ ДО ПОДПИСАНИЯ].")
        if schedule:
            document.add_paragraph("Покупатель оплачивает цену в рассрочку по датам и в суммах приложения 1. "
                                   "Проценты за предоставление рассрочки не начисляются. Это условие не отменяет "
                                   "предусмотренные законом последствия просрочки; порядок их применения требует проверки юристом.")
        else:
            document.add_paragraph(f"Полная цена подлежит единовременной оплате не позднее "
                                   f"{_display_date(clean['payment_due_date'])}. Условия раскрытия аккредитива, эскроу "
                                   "или иного выбранного механизма расчетов, если он используется: [СОГЛАСОВАТЬ].")

    document.add_heading("Передача и оформление права", level=1)
    if kind == "main":
        document.add_paragraph(f"Передача участка оформляется отдельным передаточным актом в течение "
                               f"{clean['transfer_days']} календарных дней после государственной регистрации перехода права "
                               "собственности к Покупателю. Условия доступа, фактическое состояние и передаваемые документы "
                               "указываются в акте после осмотра.")
    else:
        document.add_paragraph("Срок и порядок фактической передачи участка, содержание передаточного акта и порядок "
                               "подачи документов на регистрацию необходимо закрепить в основном договоре.")
    document.add_paragraph("Переход права собственности подлежит государственной регистрации. Порядок подачи документов, "
                           "распределение расходов и действия при приостановлении регистрации: [СОГЛАСОВАТЬ]. "
                           "Наличие подписанного проекта или платежного календаря не подтверждает регистрацию.")

    document.add_heading("Режим использования участка", level=1)
    paragraph = document.add_paragraph()
    paragraph.add_run("Участок рассматривается как полевой участок ЛПХ. Сейчас строительство запрещено. ").bold = True
    paragraph.add_run("Инициативы об изменении категории, границ или разрешенного использования не дают права "
                      "строить. Проект не обещает будущего разрешения на строительство. Фактический правовой "
                      "режим и ограничения следует сверить с актуальными документами по конкретному участку.")

    document.add_heading("Условия для юридической проверки", level=1)
    document.add_paragraph("Юрист проверяет титул продавца, историю перехода права, границы и площадь, "
                           "обременения, запреты и права третьих лиц; согласие супруга и иные необходимые согласия; "
                           "полномочия представителей, необходимость нотариальной формы и применимость специальных правил оборота земли.")
    document.add_paragraph("Залог в силу закона при рассрочке, его регистрация и снятие, последствия просрочки, "
                           "расторжения или отказа в регистрации требуют отдельного согласования. Настоящий проект "
                           "не исключает залог автоматически и не подтверждает отсутствие обременений.")
    document.add_paragraph("До подписания юрист должен увязать платежи, фактическую передачу и регистрацию права; "
                           "проверить действующее законодательство, сроки уведомлений, порядок разрешения споров "
                           "и согласовать все незаполненные условия.")

    document.add_heading("Реквизиты и подписание", level=1)
    document.add_paragraph("Полные идентификационные и адресные данные сторон, реквизиты документов и полномочий: "
                           "[ДОПОЛНИТЬ И ПРОВЕРИТЬ В ЗАЩИЩЕННОМ ПОРЯДКЕ]. Подписи в черновик не добавлены. "
                           "Окончательный текст подписывают стороны после юридической проверки и согласования условий.")
    if schedule:
        document.add_page_break()
        document.add_heading("Приложение 1 График платежей", level=1)
        document.add_paragraph(f"К проекту от {_display_date(clean['contract_date'])}. "
                               f"Общая цена: {_display_money(clean['price'])} руб. Проценты за рассрочку не начисляются.")
        if kind == "preliminary":
            document.add_paragraph("Предлагаемые условия будущего основного договора. До его заключения этот график "
                                   "не устанавливает обязанность оплаты.")
        _schedule_table(document, schedule)
        paragraph = document.add_paragraph(f"Итого: {_display_money(clean['price'])} руб. Остаток после последнего платежа: 0,00 руб.")
        paragraph.paragraph_format.space_before = Pt(9)
        document.add_paragraph("Строка 0 отражает первоначальный взнос; при сумме 0,00 руб. платеж в эту дату не требуется. "
                               "Если исходного дня нет в месяце, используется последний день этого месяца. "
                               "Разница в копейках включена в последний платеж. Правила переноса сроков, выпавших "
                               "на нерабочий день, необходимо согласовать с учетом закона и выбранного способа оплаты.")
        document.add_paragraph("График показывает плановые суммы и не подтверждает получение денег продавцом.")
    document.core_properties.title = title
    buffer = BytesIO()
    document.save(buffer)
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("xb") as file:
        file.write(buffer.getvalue())
    return output
